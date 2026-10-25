#!/usr/bin/env python3
"""Build an editable DOCX from the same controlled Markdown as the PDF.

Word pagination is renderer-dependent. The document therefore contains a real
TOC field and PAGEREF fields for figure/table lists; settings request a field
update when the file is opened. The fixed-layout PDF remains the page-keyed
reference for the change log.
"""
from __future__ import annotations

import argparse
import json
import re
from collections import defaultdict
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from PIL import Image

from build_revised_thesis import (
    ASSETS, BODY_FILES, CONTENT, DELIVERABLES, EQ_RE, FIG_RE, REF_RE,
    SOURCE_RE, TABLE_RE, chapter_from_heading, parse_table, read, scan_content,
)

BLUE = "1F4E79"
LIGHT_BLUE = "E8F1F5"
GREY = "59636B"
CONTENT_WIDTH_MM = 160
BODY_FONT = "Times New Roman"
BODY_SIZE_PT = 12


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    element = OxmlElement("w:tblHeader")
    element.set(qn("w:val"), "true")
    tr_pr.append(element)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def set_rtl(paragraph, enabled: bool = True) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    bidi = p_pr.find(qn("w:bidi"))
    if enabled and bidi is None:
        bidi = OxmlElement("w:bidi")
        p_pr.append(bidi)
    elif not enabled and bidi is not None:
        p_pr.remove(bidi)
    for run in paragraph.runs:
        r_pr = run._r.get_or_add_rPr()
        if enabled and r_pr.find(qn("w:rtl")) is None:
            r_pr.append(OxmlElement("w:rtl"))


def bookmark_name(label: str) -> str:
    return ("BM_" + label.replace("Figure", "FIG").replace("Tableau", "TAB")
            .replace(" ", "").replace(".", "_"))


def add_bookmark(paragraph, name: str, number: int) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(number))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(number))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def add_field(paragraph, instruction: str, placeholder: str = "Mettre à jour les champs") -> None:
    begin_run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    begin_run._r.append(begin)
    instruction_run = paragraph.add_run()
    instruction_text = OxmlElement("w:instrText")
    instruction_text.set(qn("xml:space"), "preserve")
    instruction_text.text = " " + instruction + " "
    instruction_run._r.append(instruction_text)
    separate_run = paragraph.add_run()
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    separate_run._r.append(separate)
    paragraph.add_run(placeholder)
    end_run = paragraph.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    end_run._r.append(end)


def add_markdown_runs(paragraph, text: str, font: str = BODY_FONT, size: float = BODY_SIZE_PT,
                      color: str | None = None, rtl: bool = False) -> None:
    # Historical call sites use Serif/Sans as semantic slots. Both now resolve
    # to the documented Algerian Master's body typeface in the editable file.
    font = BODY_FONT
    # The controlled sources use bold as their only inline Markdown convention.
    parts = re.split(r"(\*\*.*?\*\*)", text)
    for part in parts:
        if not part:
            continue
        bold = part.startswith("**") and part.endswith("**")
        value = part[2:-2] if bold else part
        run = paragraph.add_run(value)
        run.bold = bold
        run.font.name = font
        run.font.size = Pt(size)
        r_fonts = run._r.get_or_add_rPr().get_or_add_rFonts()
        for key in ("ascii", "hAnsi", "cs", "eastAsia"):
            r_fonts.set(qn(f"w:{key}"), font)
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
        if rtl:
            run._r.get_or_add_rPr().append(OxmlElement("w:rtl"))
    if rtl:
        set_rtl(paragraph, True)


def set_running_header(paragraph, text: str) -> None:
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.paragraph_format.line_spacing = 1.0
    add_markdown_runs(paragraph, text, BODY_FONT, 9, GREY)


def set_centered_page_field(paragraph) -> None:
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.line_spacing = 1.0
    add_field(paragraph, "PAGE", "1")
    for run in paragraph.runs:
        run.font.name = BODY_FONT
        run.font.size = Pt(10)


class ThesisDocx:
    def __init__(self):
        self.doc = Document()
        self.current_chapter = ""
        self.current_part = ""
        self.arabic_mode = False
        self.counts: dict[tuple[str, str], int] = defaultdict(int)
        self.captions: list[str] = []
        self.bookmark_id = 10
        self.configure()

    def configure(self) -> None:
        section = self.doc.sections[0]
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.top_margin = Mm(25)
        section.bottom_margin = Mm(25)
        section.left_margin = Mm(25)
        section.right_margin = Mm(25)
        section.header_distance = Mm(8)
        section.footer_distance = Mm(8)
        section.different_first_page_header_footer = True

        styles = self.doc.styles
        normal = styles["Normal"]
        normal.font.name = BODY_FONT
        normal.font.size = Pt(BODY_SIZE_PT)
        normal.paragraph_format.line_spacing = 1.5
        normal.paragraph_format.space_after = Pt(5)
        normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for name, size in (("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 12), ("Heading 4", 12)):
            style = styles[name]
            style.font.name = BODY_FONT
            style.font.size = Pt(size)
            style.font.bold = True
            style.font.color.rgb = RGBColor.from_string(BLUE)
            style.paragraph_format.space_before = Pt(10 if name != "Heading 1" else 0)
            style.paragraph_format.space_after = Pt(6)
            style.paragraph_format.keep_with_next = True
        styles["Heading 1"].paragraph_format.page_break_before = False

        set_running_header(section.header.paragraphs[0], "Particularités de la reconnaissance géotechnique des ponts — version révisée")
        set_centered_page_field(section.footer.paragraphs[0])

        settings = self.doc.settings.element
        update = OxmlElement("w:updateFields")
        update.set(qn("w:val"), "true")
        settings.append(update)
        self.doc.core_properties.title = "Particularités de la reconnaissance géotechnique des ponts — version révisée"
        self.doc.core_properties.author = "BERKANE Dhiyaddine"
        self.doc.core_properties.subject = "Étude bibliographique — version révisée et traçable"

    def page_break(self) -> None:
        self.doc.add_page_break()

    def start_body_section(self, heading: str) -> None:
        """Start a chapter-level section with its own running header.

        Footers remain centered PAGE fields and no numbering restart is set, so
        pagination continues while the header follows the current chapter.
        """
        section = self.doc.add_section(WD_SECTION.NEW_PAGE)
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.top_margin = Mm(25)
        section.bottom_margin = Mm(25)
        section.left_margin = Mm(25)
        section.right_margin = Mm(25)
        section.header_distance = Mm(8)
        section.footer_distance = Mm(8)
        section.different_first_page_header_footer = True
        for part in (section.header, section.footer, section.first_page_header, section.first_page_footer):
            part.is_linked_to_previous = False
        set_running_header(section.header.paragraphs[0], heading)
        section.first_page_header.paragraphs[0].clear()
        set_centered_page_field(section.footer.paragraphs[0])
        set_centered_page_field(section.first_page_footer.paragraphs[0])

    def add_cover(self) -> None:
        for text in ("République Algérienne Démocratique et Populaire",
                     "Ministère de l’Enseignement Supérieur et de la Recherche Scientifique"):
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_markdown_runs(p, text, "Serif", 12)
            for run in p.runs:
                run.bold = True
        logos = self.doc.add_table(rows=1, cols=2)
        logos.alignment = WD_TABLE_ALIGNMENT.CENTER
        for cell, filename in zip(logos.rows[0].cells, ("cover_logo_1.png", "cover_logo_2.jpeg")):
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(ASSETS / filename), width=Mm(20))
        for text, size in (("Université Larbi Tébessi – Tébessa", 11.5),
                           ("Faculté des Sciences Exactes, des Sciences de la Nature et de la Vie", 11),
                           ("Département des Sciences de la Terre et de l’Univers", 11)):
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_markdown_runs(p, text, "Serif", size)
            for run in p.runs:
                run.bold = True
        self.doc.add_paragraph()
        for text, size in (("Mémoire présenté en vue de l’obtention du diplôme de", 11),
                           ("MASTER", 14), ("Filière : Géologie", 11),
                           ("Option : Géologie de l’Ingénieur et Géotechnique", 11)):
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_markdown_runs(p, text, "Serif", size)
            for run in p.runs:
                run.bold = text == "MASTER"
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(10)
        add_markdown_runs(p, "**Particularités de la reconnaissance géotechnique des ponts**\n(Étude bibliographique)", BODY_FONT, 16)
        p_pr = p._p.get_or_add_pPr()
        borders = OxmlElement("w:pBdr")
        for side in ("top", "left", "bottom", "right"):
            edge = OxmlElement(f"w:{side}")
            edge.set(qn("w:val"), "single")
            edge.set(qn("w:sz"), "8")
            edge.set(qn("w:color"), BLUE)
            borders.append(edge)
        p_pr.append(borders)
        for text in ("Présenté par : BERKANE Dhiyaddine", "Devant le jury :"):
            p = self.doc.add_paragraph()
            add_markdown_runs(p, f"**{text}**", "Serif", 10.5)
        jury = [("Pr BAALI Fathi", "Professeur", "U.L.T. Tébessa", "Président"),
                ("Pr FEHDI Chems eddine", "Professeur", "U.L.T. Tébessa", "Examinateur"),
                ("Dr MOUICI Ridha", "M.C. « B »", "U.L.T. Tébessa", "Promoteur, Rapporteur"),
                ("Dr AOUN Mounira", "M.C. « B »", "U.L.T. Tébessa", "Co-promoteur")]
        table = self.doc.add_table(rows=0, cols=4)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for values in jury:
            cells = table.add_row().cells
            for cell, value in zip(cells, values):
                cell.text = value
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.name = BODY_FONT
                        run.font.size = Pt(8.5)
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_markdown_runs(p, "**Session Juin 2026**", "Serif", 10)
        # Cover typography is compact and single-spaced; the 1.5 rule applies
        # to the manuscript body, not to the institutional identification block.
        for paragraph in self.doc.paragraphs:
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.paragraph_format.space_after = Pt(2)
        for cover_table in self.doc.tables:
            for row in cover_table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.line_spacing = 1.0
                        paragraph.paragraph_format.space_after = Pt(0)

    def add_major_heading(self, text: str, body: bool) -> None:
        if body:
            self.start_body_section(text)
        else:
            self.page_break()
        self.current_part = text
        chapter = chapter_from_heading(text)
        if chapter:
            self.current_chapter = chapter
        elif text.startswith(("INTRODUCTION", "CONCLUSION", "BIBLIOGRAPHIE")):
            self.current_chapter = ""
        self.arabic_mode = text.strip() == "الملخص"
        p = self.doc.add_paragraph(style="Heading 1" if body else None)
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if self.arabic_mode else WD_ALIGN_PARAGRAPH.LEFT
        add_markdown_runs(p, text, "Sans", 14, BLUE, self.arabic_mode)
        for run in p.runs:
            run.bold = True
        p_pr = p._p.get_or_add_pPr()
        borders = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "8")
        bottom.set(qn("w:color"), BLUE)
        borders.append(bottom)
        p_pr.append(borders)

    def add_heading(self, text: str, level: int) -> None:
        p = self.doc.add_paragraph(style=f"Heading {min(level, 4)}")
        add_markdown_runs(p, text, BODY_FONT, BODY_SIZE_PT, BLUE if level == 2 else "233741")
        for run in p.runs:
            run.bold = True

    def add_paragraph(self, text: str) -> None:
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if self.arabic_mode else WD_ALIGN_PARAGRAPH.JUSTIFY
        add_markdown_runs(p, text, BODY_FONT, BODY_SIZE_PT, rtl=self.arabic_mode)

    def add_bullet(self, text: str, ordered: str | None = None) -> None:
        p = self.doc.add_paragraph()
        p.paragraph_format.left_indent = Mm(4)
        p.paragraph_format.first_line_indent = Mm(-3)
        p.paragraph_format.space_after = Pt(2)
        add_markdown_runs(p, f"{ordered + '. ' if ordered else '• '}{text}", BODY_FONT, BODY_SIZE_PT)

    def add_callout(self, text: str) -> None:
        p = self.doc.add_paragraph()
        p.paragraph_format.left_indent = Mm(2)
        p.paragraph_format.right_indent = Mm(2)
        add_markdown_runs(p, text, BODY_FONT, BODY_SIZE_PT)
        p_pr = p._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), LIGHT_BLUE)
        p_pr.append(shd)

    def add_equation(self, text: str) -> None:
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_markdown_runs(p, text, BODY_FONT, BODY_SIZE_PT)

    def add_source(self, text: str) -> None:
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        add_markdown_runs(p, "Source : " + text, BODY_FONT, 10, GREY)

    def add_reference(self, text: str) -> None:
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.left_indent = Mm(5)
        p.paragraph_format.first_line_indent = Mm(-5)
        p.paragraph_format.space_after = Pt(4)
        add_markdown_runs(p, text, BODY_FONT, BODY_SIZE_PT)

    def add_figure(self, asset: str, title: str) -> None:
        if not self.current_chapter:
            raise ValueError(f"Figure outside chapter: {asset}")
        key = ("Figure", self.current_chapter)
        self.counts[key] += 1
        label = f"Figure {self.current_chapter}.{self.counts[key]}"
        path = ASSETS / "figures" / f"{asset}.png"
        with Image.open(path) as image:
            aspect = image.height / image.width
        width = 160
        if width * aspect > 145:
            width = 145 / aspect
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(str(path), width=Mm(width))
        caption = self.doc.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_markdown_runs(caption, f"**{label} — {title}**", BODY_FONT, 10)
        add_bookmark(caption, bookmark_name(label), self.bookmark_id)
        self.bookmark_id += 1
        self.captions.append(label)

    def add_table(self, rows: list[list[str]], caption: str | None, weights: str | None = None) -> None:
        if not rows:
            return
        if any(len(row) != len(rows[0]) for row in rows):
            raise ValueError(f"Malformed table near {caption}")
        if caption:
            if not self.current_chapter:
                raise ValueError(f"Numbered table outside chapter: {caption}")
            key = ("Tableau", self.current_chapter)
            self.counts[key] += 1
            label = f"Tableau {self.current_chapter}.{self.counts[key]}"
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            add_markdown_runs(p, f"**{label} — {caption}**", BODY_FONT, 10)
            add_bookmark(p, bookmark_name(label), self.bookmark_id)
            self.bookmark_id += 1
            self.captions.append(label)
        table = self.doc.add_table(rows=len(rows), cols=len(rows[0]))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        table.style = "Table Grid"
        values = [float(v) for v in weights.split(",")] if weights else [1] * len(rows[0])
        if len(values) != len(rows[0]):
            raise ValueError(f"Column weight mismatch: {caption}")
        total = sum(values)
        for r_index, (row, values_row) in enumerate(zip(table.rows, rows)):
            prevent_row_split(row)
            if r_index == 0:
                set_repeat_table_header(row)
            for c_index, (cell, text) in enumerate(zip(row.cells, values_row)):
                cell.width = Mm(CONTENT_WIDTH_MM * values[c_index] / total)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                if r_index == 0:
                    set_cell_shading(cell, BLUE)
                elif r_index % 2 == 0:
                    set_cell_shading(cell, "F3F6F8")
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(1.5)
                add_markdown_runs(p, text, BODY_FONT, 8.5, "FFFFFF" if r_index == 0 else None)
                if r_index == 0:
                    for run in p.runs:
                        run.bold = True
        self.doc.add_paragraph().paragraph_format.space_after = Pt(1)

    def add_caption_list(self, kind: str, templates) -> None:
        self.page_break()
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_markdown_runs(p, f"**LISTE DES {kind.upper()}S**", "Sans", 14, BLUE)
        note = self.doc.add_paragraph()
        add_markdown_runs(note, "Numéros de page actualisables dans Word (champs PAGEREF).", "Serif", 8, GREY)
        for item in templates:
            p = self.doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.tab_stops.add_tab_stop(Mm(CONTENT_WIDTH_MM), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
            add_markdown_runs(p, f"{item.label} — {item.title}\t", "Serif", 8.8)
            add_field(p, f"PAGEREF {bookmark_name(item.label)} \\h", "—")

    def add_toc(self) -> None:
        self.page_break()
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_markdown_runs(p, "**TABLE DES MATIÈRES**", "Sans", 14, BLUE)
        toc = self.doc.add_paragraph()
        add_field(toc, 'TOC \\o "1-2" \\h \\z \\u', "Dans Word : clic droit, puis « Mettre à jour les champs ».")
        note = self.doc.add_paragraph()
        add_markdown_runs(note, "La table est un champ Word dynamique ; le PDF fournit la pagination fixe validée.", "Serif", 8, GREY)

    def render(self, text: str, body: bool = True) -> None:
        lines = text.splitlines()
        paragraph: list[str] = []
        pending_table: tuple[str, str | None] | None = None

        def flush() -> None:
            nonlocal paragraph
            if paragraph:
                value = " ".join(part.strip() for part in paragraph).strip()
                if value:
                    if self.current_part.startswith("BIBLIOGRAPHIE"):
                        self.add_reference(value)
                    else:
                        self.add_paragraph(value)
                paragraph = []

        i = 0
        while i < len(lines):
            raw = lines[i]
            line = raw.strip()
            if not line:
                flush()
                i += 1
                continue
            if line.startswith("|"):
                flush()
                rows, i = parse_table(lines, i)
                caption, weights = pending_table or (None, None)
                self.add_table(rows, caption, weights)
                pending_table = None
                continue
            if line.startswith("# "):
                flush()
                self.add_major_heading(line[2:].strip(), body)
            elif line.startswith("## "):
                flush()
                self.add_heading(line[3:].strip(), 2)
            elif line.startswith("### "):
                flush()
                self.add_heading(line[4:].strip(), 3)
            elif line.startswith("#### "):
                flush()
                self.add_heading(line[5:].strip(), 4)
            elif line == "[[PAGEBREAK]]":
                flush()
                self.page_break()
            else:
                fig = FIG_RE.match(line)
                tab = TABLE_RE.match(line)
                eq = EQ_RE.match(line)
                source = SOURCE_RE.match(line)
                ref = REF_RE.match(line)
                if fig:
                    flush()
                    self.add_figure(fig.group(1), fig.group(2).strip())
                elif tab:
                    flush()
                    pending_table = (tab.group(1).strip(), tab.group(2))
                elif eq:
                    flush()
                    self.add_equation(eq.group(1).strip())
                elif source:
                    flush()
                    self.add_source(source.group(1).strip())
                elif ref:
                    flush()
                    self.add_reference(ref.group(1).strip())
                elif line.startswith("> "):
                    flush()
                    self.add_callout(line[2:].strip())
                elif re.match(r"^[-*] ", line):
                    flush()
                    self.add_bullet(line[2:].strip())
                elif re.match(r"^\d+\. ", line):
                    flush()
                    match = re.match(r"^(\d+)\.\s+(.+)$", line)
                    self.add_bullet(match.group(2), match.group(1))
                else:
                    paragraph.append(line)
            i += 1
        flush()
        if pending_table:
            raise ValueError(f"Table caption without table: {pending_table}")

    def save(self, output: Path) -> dict[str, object]:
        self.doc.save(output)
        # Reopen immediately to catch malformed package/XML output.
        reopened = Document(output)
        report = {
            "output": str(output), "paragraphs": len(reopened.paragraphs),
            "sections": len(reopened.sections),
            "tables": len(reopened.tables), "inline_shapes": len(reopened.inline_shapes),
            "numbered_captions": len(self.captions), "unique_captions": len(set(self.captions)),
            "bytes": output.stat().st_size,
            "field_update_on_open": True,
            "formatting_profile": {
                "paper": "A4", "margins_mm": {"top": 25, "right": 25, "bottom": 25, "left": 25},
                "body_font": BODY_FONT, "body_size_pt": BODY_SIZE_PT,
                "body_line_spacing": 1.5, "paragraph_alignment": "justified",
                "chapter_titles": "Times New Roman 14 pt bold",
                "section_titles": "Times New Roman 12 pt bold",
                "figure_captions": "centered below, 10 pt", "table_captions": "centered above, 10 pt",
                "sources": "right aligned, 10 pt", "pagination": "centered PAGE field in footer",
            },
        }
        if report["inline_shapes"] != 23:  # two cover logos + 21 technical figures
            raise ValueError(f"Unexpected DOCX image count: {report}")
        if len(self.captions) != 67 or len(set(self.captions)) != 67:
            raise ValueError(f"Caption mismatch: {report}")
        return report


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, default=DELIVERABLES / "These_revisee_BERKANE_Dhiyaddine.docx")
    parser.add_argument("--report", type=Path, default=Path(__file__).resolve().parents[1] / "audits" / "docx_build_report.json")
    args = parser.parse_args()
    args.output.parent.mkdir(parents=True, exist_ok=True)
    scan = scan_content()
    thesis = ThesisDocx()
    thesis.add_cover()
    thesis.render(read(CONTENT / "00_front.md"), body=False)
    thesis.arabic_mode = False
    thesis.add_caption_list("Figure", [item for item in scan.captions if item.kind == "Figure"])
    thesis.add_caption_list("Tableau", [item for item in scan.captions if item.kind == "Tableau"])
    thesis.render(read(CONTENT / "00_abbreviations.md"), body=False)
    thesis.add_toc()
    for filename in BODY_FILES:
        thesis.render(read(CONTENT / filename), body=True)
    report = thesis.save(args.output)
    args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
