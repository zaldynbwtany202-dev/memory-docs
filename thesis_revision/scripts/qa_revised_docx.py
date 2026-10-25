#!/usr/bin/env python3
"""Reopen and audit the editable DOCX package and its controlled content."""
from __future__ import annotations

import hashlib
import json
import re
import zipfile
from collections import Counter
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

from build_revised_thesis import ASSETS, BODY_FILES, CONTENT, ROOT, read, scan_content

DOCX = ROOT / "deliverables" / "These_revisee_BERKANE_Dhiyaddine.docx"
BUILD_REPORT = ROOT / "audits" / "docx_build_report.json"
OUT = ROOT / "audits" / "docx_content_audit.json"


def sha256_bytes(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


build = json.loads(BUILD_REPORT.read_text(encoding="utf-8"))
scan = scan_content()
with zipfile.ZipFile(DOCX) as package:
    corrupt_member = package.testzip()
    document_xml = package.read("word/document.xml")
    settings_xml = package.read("word/settings.xml")
    media = {name: package.read(name) for name in package.namelist() if name.startswith("word/media/") and not name.endswith("/")}

doc = Document(DOCX)
body = doc.element.body
all_text = "\n".join(body.itertext())
ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
bookmark_names = [node.get(qn("w:name"), "") for node in body.xpath(".//w:bookmarkStart")]
instructions = ["".join(node.itertext()).strip() for node in body.xpath(".//w:instrText")]
bidi_paragraphs = len(body.xpath(".//w:p[w:pPr/w:bidi]"))
update_nodes = doc.settings.element.xpath(".//w:updateFields")
update_fields = bool(update_nodes) and update_nodes[0].get(qn("w:val")) in {"true", "1", "on"}

expected_caption_bookmarks = {
    "BM_" + item.label.replace("Figure", "FIG").replace("Tableau", "TAB")
    .replace(" ", "").replace(".", "_")
    for item in scan.captions
}
expected_media = [ASSETS / "cover_logo_1.png", ASSETS / "cover_logo_2.jpeg"] + [ASSETS / "figures" / f"{item.asset}.png" for item in scan.captions if item.kind == "Figure"]
# The scan contains captions in reading order; each figure key is unique.
expected_hashes = Counter(sha256_file(path) for path in expected_media)
embedded_hashes = Counter(sha256_bytes(value) for value in media.values())

jury_names = ["Pr BAALI Fathi", "Pr FEHDI Chems eddine", "Dr MOUICI Ridha", "Dr AOUN Mounira"]
expected_chapter_headers = [
    line.strip()[2:].strip()
    for filename in BODY_FILES
    for line in read(CONTENT / filename).splitlines()
    if line.strip().startswith("# ")
]
chapter_headers = [section.header.paragraphs[0].text for section in doc.sections[1:]]
footer_fields_ok = all(
    any("PAGE" in "".join(node.itertext()) for node in section.footer._element.xpath(".//w:instrText"))
    and section.footer.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
    for section in doc.sections
)
first_page_footer_fields_ok = all(
    any("PAGE" in "".join(node.itertext()) for node in section.first_page_footer._element.xpath(".//w:instrText"))
    and section.first_page_footer.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
    for section in doc.sections[1:]
)
normal = doc.styles["Normal"]
heading_styles_ok = (
    doc.styles["Heading 1"].font.name == "Times New Roman"
    and abs(doc.styles["Heading 1"].font.size.pt - 14) < 0.01
    and all(
        doc.styles[f"Heading {level}"].font.name == "Times New Roman"
        and abs(doc.styles[f"Heading {level}"].font.size.pt - 12) < 0.01
        and doc.styles[f"Heading {level}"].font.bold
        for level in (2, 3, 4)
    )
)
caption_paragraphs = [p for p in doc.paragraphs if p._p.xpath("./w:bookmarkStart")]
source_paragraphs = [p for p in doc.paragraphs if p.text.startswith("Source :")]
caption_format_ok = (
    len(caption_paragraphs) == len(scan.captions)
    and all(p.alignment == WD_ALIGN_PARAGRAPH.CENTER for p in caption_paragraphs)
    and all(all(run.font.size is not None and abs(run.font.size.pt - 10) < 0.01 for run in p.runs if run.text) for p in caption_paragraphs)
)
source_format_ok = (
    len(source_paragraphs) >= sum(item.kind == "Figure" for item in scan.captions)
    and all(p.alignment == WD_ALIGN_PARAGRAPH.RIGHT for p in source_paragraphs)
    and all(all(run.font.size is not None and abs(run.font.size.pt - 10) < 0.01 for run in p.runs if run.text) for p in source_paragraphs)
)
profile = build.get("formatting_profile", {})
checks = {
    "zip_package_integrity": corrupt_member is None,
    "python_docx_reopen": len(doc.paragraphs) == build["paragraphs"],
    "a4_page_setup": all(abs(section.page_width.mm - 210) < 0.1 and abs(section.page_height.mm - 297) < 0.1 for section in doc.sections),
    "documented_formatting_profile": (
        profile.get("paper") == "A4"
        and profile.get("margins_mm") == {"top": 25, "right": 25, "bottom": 25, "left": 25}
        and profile.get("body_font") == "Times New Roman"
        and profile.get("body_size_pt") == 12
        and profile.get("body_line_spacing") == 1.5
    ),
    "margins_25mm_all_sections": all(
        all(abs(value.mm - 25) < 0.1 for value in (section.top_margin, section.right_margin, section.bottom_margin, section.left_margin))
        for section in doc.sections
    ),
    "times_new_roman_12_and_1_5_spacing": (
        normal.font.name == "Times New Roman"
        and normal.font.size is not None and abs(normal.font.size.pt - 12) < 0.01
        and normal.paragraph_format.line_spacing == 1.5
        and normal.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
        and heading_styles_ok
    ),
    "chapter_headers_and_continuous_page_fields": (
        len(doc.sections) == len(expected_chapter_headers) + 1
        and chapter_headers == expected_chapter_headers
        and footer_fields_ok and first_page_footer_fields_ok
    ),
    "caption_and_source_typography": caption_format_ok and source_format_ok,
    "cover_jury_names": all(name in all_text for name in jury_names),
    "no_placeholder_names": not any(token in all_text for token in ["NOM PRÉNOM", "NOM ET PRÉNOM", "Nom du jury", "Nom du promoteur"]),
    "arabic_abstract": len(re.findall(r"[\u0600-\u06ff]", all_text)) > 1000 and "تتناول" in all_text and "الكلمات" in all_text,
    "roman_example": "Alcántara" in all_text and "CORTÉS PÉREZ" in all_text,
    "luzancy_source": "Luzancy" in all_text and "XERCAVINS" in all_text,
    "bibliography_65th_reference": "Liquefaction Resistance of Soils: Summary Report" in all_text,
    "expected_table_count": len(doc.tables) == build["tables"],
    "expected_inline_shape_count": len(doc.inline_shapes) == build["inline_shapes"] == 2 + sum(item.kind == "Figure" for item in scan.captions),
    "caption_bookmarks_complete_and_unique": len(bookmark_names) == len(scan.captions) and set(bookmark_names) == expected_caption_bookmarks and len(set(bookmark_names)) == len(bookmark_names),
    "pageref_fields_match_captions": sum(value.startswith("PAGEREF ") for value in instructions) == len(scan.captions),
    "single_toc_field": sum(value.startswith("TOC ") for value in instructions) == 1,
    "update_fields_on_open": update_fields and b"<w:updateFields" in settings_xml,
    "embedded_media_match_current_assets": len(media) == len(expected_media) and embedded_hashes == expected_hashes,
    "no_unresolved_markers_or_known_source_defects": not any(token in all_text for token in ["[[", "]]", "浸水", "214 références", "Error!", "PENDING", "�"]),
    "six_rtl_paragraphs": bidi_paragraphs == 6,
}
report = {
    "file": str(DOCX),
    "bytes": DOCX.stat().st_size,
    "sha256": sha256_file(DOCX),
    "paragraphs": len(doc.paragraphs),
    "sections": len(doc.sections),
    "tables": len(doc.tables),
    "inline_shapes": len(doc.inline_shapes),
    "bidi_paragraphs": bidi_paragraphs,
    "bookmarks": len(bookmark_names),
    "pageref_fields": sum(value.startswith("PAGEREF ") for value in instructions),
    "toc_fields": sum(value.startswith("TOC ") for value in instructions),
    "media_files": len(media),
    "formatting_evidence": {
        "profile": profile,
        "normal_font": normal.font.name,
        "normal_size_pt": normal.font.size.pt if normal.font.size else None,
        "normal_line_spacing": normal.paragraph_format.line_spacing,
        "chapter_headers": chapter_headers,
        "caption_paragraphs_checked": len(caption_paragraphs),
        "source_paragraphs_checked": len(source_paragraphs),
    },
    "checks": checks,
    "details": {
        "corrupt_member": corrupt_member,
        "missing_bookmarks": sorted(expected_caption_bookmarks - set(bookmark_names)),
        "unexpected_bookmarks": sorted(set(bookmark_names) - expected_caption_bookmarks),
        "missing_media_hashes": sorted((expected_hashes - embedded_hashes).elements()),
        "unexpected_media_hashes": sorted((embedded_hashes - expected_hashes).elements()),
    },
    "rendering_note": "Package, structure and content are validated. Word-specific pagination is renderer-dependent and has not been visually rendered in this environment.",
    "status": "pass" if all(checks.values()) else "fail",
}
OUT.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps({key: value for key, value in report.items() if key != "details"}, ensure_ascii=False, indent=2))
if report["status"] != "pass":
    print(json.dumps(report["details"], ensure_ascii=False, indent=2))
raise SystemExit(0 if report["status"] == "pass" else 2)
